"""2026-09-22 deep review, slice I — 08_Documentation_Word (I-09, I-10, I-16, I-17, I-24).

Drives the real collector -> Word generator -> snapshot chain with a fake
connection (INFO-shaped rows with integer enums, TMSL-shaped snapshot) and
real python-docx output under tmp_path.
"""

from __future__ import annotations

import copy
import json
import os
from types import SimpleNamespace

import pytest

pytest.importorskip("docx")

from docx import Document  # noqa: E402

import server.handlers.documentation_handler as handler  # noqa: E402
from core.analysis.bpa_analyzer import BPAAnalyzer  # noqa: E402
from core.documentation.data_collector import collect_model_documentation  # noqa: E402
from core.documentation.utils import (  # noqa: E402
    REPO_ROOT,
    ensure_dir,
    resolve_doc_target,
    snapshot_path_for,
)
from core.orchestration.documentation_orchestrator import DocumentationOrchestrator  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

_TMSL = {
    "name": "Model",
    "tables": [
        {
            "name": "Sales",
            "columns": [{"name": "Amount", "dataType": "double", "sourceColumn": "Amount"}],
            "measures": [{"name": "Bad", "expression": "SUMX ( FILTER ( Sales, 1 ), 1 )"}],
            "partitions": [
                {"name": "Sales", "source": {"type": "m", "expression": ["let", " S = 1", "in S"]}}
            ],
        },
        {
            "name": "Time Intel",
            "calculationGroup": {
                "precedence": 10,
                "calculationItems": [
                    {"name": "CY", "expression": "SELECTEDMEASURE()", "ordinal": 0}
                ],
            },
            "columns": [{"name": "Name", "dataType": "string"}],
        },
    ],
    "roles": [
        {
            "name": "Region",
            "members": [{"memberName": "a@contoso.example"}],
            "tablePermissions": [{"name": "Sales", "filterExpression": "[Amount] > 0"}],
        }
    ],
    "functions": [{"name": "Fx.Double", "expression": "(x) => x * 2"}],
}


class _Exec:
    def __init__(self, measures):
        self.measures = measures
        self.tmsl = copy.deepcopy(_TMSL)

    def execute_info_query(self, fn, **kw):
        rows = {
            "TABLES": [
                {"ID": 1, "Name": "Sales", "Description": "Fact\x0b table"},
                {"ID": 2, "Name": "Time Intel"},
            ],
            "COLUMNS": [
                {
                    "Table": "Sales",
                    "ExplicitName": "Zeta",
                    "Name": "Zeta",
                    "ExplicitDataType": 10,
                    "DataType": 10,
                    "Type": 1,
                },
                {
                    "Table": "Sales",
                    "ExplicitName": "Alpha",
                    "Name": "Alpha",
                    "ExplicitDataType": 2,
                    "DataType": 2,
                    "Type": 2,
                    "Description": "calc\x00col",
                },
                {
                    "Table": "Sales",
                    "ExplicitName": "RowNumber-1",
                    "Name": "RowNumber-1",
                    "ExplicitDataType": 6,
                    "Type": 3,
                },
            ],
            "MEASURES": self.measures,
            "RELATIONSHIPS": [
                {
                    "FromTable": "Sales",
                    "FromColumn": "Alpha",
                    "ToTable": "Time Intel",
                    "ToColumn": "Name",
                    "FromCardinality": 2,
                    "ToCardinality": 1,
                    "CrossFilteringBehavior": 2,
                    "IsActive": True,
                }
            ],
        }.get(fn, [])
        return {"success": True, "rows": [dict(r) for r in rows]}

    def get_tmsl_definition(self):
        return {"success": True, "tmsl": json.dumps(self.tmsl)}

    def _get_database_name(self):
        return "0f3a1c2e-1111-2222-3333-444455556666"


def _state(measures=None):
    measures = (
        measures
        if measures is not None
        else [
            {"Table": "Sales", "Name": "Total", "Expression": "SUM ( Sales[Amount] )"},
            {"Table": "Sales", "Name": "Avg", "Expression": "AVERAGE ( Sales[Amount] )"},
        ]
    )
    return SimpleNamespace(
        is_connected=lambda: True,
        query_executor=_Exec(measures),
        model_exporter=None,
        bpa_analyzer=BPAAnalyzer(os.path.join(ROOT, "config", "bpa_rules_comprehensive.json")),
        model_validator=None,
        connection_manager=SimpleNamespace(
            get_instance_info=lambda: {"display_name": "Contoso Sales.pbix"}
        ),
    )


# ── utils (I-24 / I-09) ──────────────────────────────────────────────────────


def test_default_exports_dir_is_repo_exports_not_core_exports():
    d = ensure_dir(None)
    assert os.path.normpath(d) == os.path.normpath(os.path.join(REPO_ROOT, "exports", "docs"))
    assert os.path.isdir(os.path.join(REPO_ROOT, "core"))


def test_docx_output_path_is_a_file_not_a_directory(tmp_path):
    directory, file = resolve_doc_target(str(tmp_path / "MyModel.docx"))
    assert file.endswith("MyModel.docx") and not os.path.isdir(file)
    assert os.path.normpath(directory) == os.path.normpath(str(tmp_path))
    assert snapshot_path_for(file).endswith("MyModel.snapshot.json")


# ── collector (I-10 / I-16) ──────────────────────────────────────────────────


def test_collector_runs_real_bpa_and_maps_info_fields():
    ctx = collect_model_documentation(_state())
    bp = ctx["best_practices"]
    assert bp["source"] == "bpa"
    assert "AVOID_SUMX_FILTER" in {v["rule_id"] for v in bp["violations"]}
    assert all(v["severity"] in ("INFO", "WARNING", "ERROR") for v in bp["violations"])
    sales = next(t for t in ctx["tables"] if t["name"] == "Sales")
    assert [c["name"] for c in sales["columns"]] == ["Alpha", "Zeta"]  # sorted, rowNumber dropped
    assert (
        sales["columns"][0]["data_type"] == "String" and sales["columns"][0]["type"] == "Calculated"
    )
    assert sales["columns"][1]["data_type"] == "Decimal"
    assert [m["name"] for m in sales["measures"]] == ["Avg", "Total"]
    rel = ctx["relationships"][0]
    assert rel["cardinality"] == "Many:One" and rel["direction"] == "BothDirections"
    assert ctx["database_name"] == "Contoso Sales.pbix"


def test_collector_extras_from_tmsl():
    ctx = collect_model_documentation(_state())
    assert ctx["calculation_groups"][0]["precedence"] == 10
    assert ctx["roles"][0]["table_permissions"][0]["filter_expression"] == "[Amount] > 0"
    assert ctx["power_query"][0]["table"] == "Sales"
    assert ctx["functions"][0]["name"] == "Fx.Double"


# ── generate / update (I-09, I-17) ───────────────────────────────────────────


def _texts(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            parts.extend(c.text for c in row.cells)
    return "\n".join(parts)


def test_generate_writes_exact_docx_and_sibling_snapshot(tmp_path):
    target = str(tmp_path / "MyModel.docx")
    res = DocumentationOrchestrator(config={}).generate_word_documentation(_state(), target)
    assert res["success"], res
    assert os.path.normpath(res["doc_path"]) == os.path.normpath(target) and os.path.isfile(target)
    assert os.path.isfile(res["snapshot_path"]) and res["snapshot_path"].endswith(
        "MyModel.snapshot.json"
    )
    text = _texts(target)
    assert "Calculation Groups" in text and "Security (RLS)" in text and "Fx.Double" in text
    assert "Fact table" in text  # \x0b stripped instead of crashing python-docx


def test_update_finds_the_snapshot_next_to_the_docx_and_logs_changes(tmp_path):
    target = str(tmp_path / "MyModel.docx")
    orch = DocumentationOrchestrator(config={})
    assert orch.generate_word_documentation(_state(), target)["success"]
    changed = _state(
        [
            {"Table": "Sales", "Name": "Total", "Expression": "SUM ( Sales[Amount] ) * 2"},
            {"Table": "Sales", "Name": "Margin", "Expression": "1"},
        ]
    )
    res = orch.update_word_documentation(changed, input_path=target)
    assert res["success"], res
    assert res["in_place"] is True and os.path.normpath(res["doc_path"]) == os.path.normpath(target)
    assert os.path.normpath(res["previous_snapshot"]) == os.path.normpath(snapshot_path_for(target))
    cs = res["change_summary"]
    assert cs["changes_detected"] is True and not cs.get("baseline_missing")
    assert cs["measures"]["added"] == ["Sales::Margin"]
    assert cs["measures"]["removed"] == ["Sales::Avg"]
    assert cs["measures"]["updated"] == ["Sales::Total"]
    assert "Change Log" in _texts(target)


def test_update_legacy_positional_call_treats_docx_as_input(tmp_path):
    target = str(tmp_path / "Legacy.docx")
    orch = DocumentationOrchestrator(config={})
    orch.generate_word_documentation(_state(), target)
    res = orch.update_word_documentation(_state(), target, None)  # old handler call shape
    assert res["success"] and res["previous_snapshot"] and res["in_place"]
    assert res["change_summary"]["changes_detected"] is False


def test_update_without_snapshot_says_baseline(tmp_path):
    target = tmp_path / "Orphan.docx"
    Document().save(str(target))
    res = DocumentationOrchestrator(config={}).update_word_documentation(
        _state(), input_path=str(target)
    )
    assert res["success"] and res["change_summary"].get("baseline_missing") is True
    assert "baseline" in res["message"]


def test_locked_output_returns_structured_error(tmp_path, monkeypatch):
    import docx.document

    def boom(self, path):
        raise PermissionError("locked")

    monkeypatch.setattr(docx.document.Document, "save", boom)
    res = DocumentationOrchestrator(config={}).generate_word_documentation(
        _state(), str(tmp_path / "L.docx")
    )
    assert res["success"] is False and res["error_type"] == "file_locked"


# ── handler ──────────────────────────────────────────────────────────────────


def test_handler_update_requires_input_and_rejects_unknown_mode(monkeypatch):
    cs = SimpleNamespace(
        is_connected=lambda: True,
        agent_policy=SimpleNamespace(documentation_orch=DocumentationOrchestrator(config={})),
    )
    monkeypatch.setattr(handler, "connection_state", cs)
    assert (
        handler.handle_documentation_word({"mode": "update"})["error_type"] == "missing_input_path"
    )
    assert handler.handle_documentation_word({"mode": "bogus"})["error_type"] == "invalid_input"


def test_handler_generate_to_docx_path_inside_exports(monkeypatch):
    target = os.path.join(REPO_ROOT, "exports", "docs", "wp10_handler_test.docx")
    st = _state()
    st.agent_policy = SimpleNamespace(documentation_orch=DocumentationOrchestrator(config={}))
    monkeypatch.setattr(handler, "connection_state", st)
    try:
        res = handler.handle_documentation_word({"mode": "generate", "output_path": target})
        assert res["success"], res
        assert os.path.isfile(target) and not os.path.isdir(target)
    finally:
        for p in (target, snapshot_path_for(target)):
            if os.path.exists(p):
                os.remove(p)
